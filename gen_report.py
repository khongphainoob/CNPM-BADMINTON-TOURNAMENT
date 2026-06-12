# -*- coding: utf-8 -*-
"""Generate DOCX report: Module THAM GIA & PHÂN CÔNG LỊCH ĐẤU UI."""
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

NAVY = RGBColor(0x1F, 0x3A, 0x5F)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

doc = Document()

base = doc.styles['Normal']
base.font.name = 'Times New Roman'
base.font.size = Pt(12)
base._element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')


def set_cell_bg(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:fill'), hex_color)
    tcPr.append(shd)


def set_cell_text(cell, text, bold=False, color=None, align='left', size=11):
    cell.text = ''
    p = cell.paragraphs[0]
    p.alignment = {'left': WD_ALIGN_PARAGRAPH.LEFT,
                   'center': WD_ALIGN_PARAGRAPH.CENTER}[align]
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = color


def add_table(headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = t.rows[0].cells
    for i, h in enumerate(headers):
        set_cell_bg(hdr[i], '1F3A5F')
        set_cell_text(hdr[i], h, bold=True, color=WHITE, align='center')
    for r in rows:
        cells = t.add_row().cells
        for i, val in enumerate(r):
            align = 'center' if (widths and i == 0) else 'left'
            set_cell_text(cells[i], str(val), align=align)
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Inches(w)
    return t


def h_title(text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(16)
    run.font.color.rgb = NAVY
    run.font.name = 'Times New Roman'
    return p


def h_section(text):
    p = doc.add_paragraph()
    run = p.add_run('●  ' + text)
    run.bold = True
    run.font.size = Pt(14)
    run.font.color.rgb = NAVY
    return p


def h_sub(text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.italic = True
    run.font.size = Pt(12)
    run.font.color.rgb = NAVY
    return p


def body(text):
    p = doc.add_paragraph(text)
    p.paragraph_format.space_after = Pt(6)
    return p


def img_placeholder(caption):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run('[ Chụp ảnh màn hình UI — dán vào đây ]')
    run.italic = True
    run.font.color.rgb = RGBColor(0x99, 0x99, 0x99)
    run.font.size = Pt(11)
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cr = cap.add_run(caption)
    cr.italic = True
    cr.font.size = Pt(10)
    cap.paragraph_format.space_after = Pt(10)


# ============ COVER ============
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = title.add_run('BÁO CÁO THIẾT KẾ GIAO DIỆN (UI)')
r.bold = True
r.font.size = Pt(20)
r.font.color.rgb = NAVY

sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = sub.add_run('MODULE: THAM GIA & PHÂN CÔNG LỊCH ĐẤU')
r.bold = True
r.font.size = Pt(16)
r.font.color.rgb = NAVY

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run('Hệ thống Quản lý Giải đấu Cầu lông — ShuttleOps').italic = True
doc.add_paragraph()

# ============ DANH SÁCH MÀN HÌNH ============
h_section('Danh sách các màn hình')
add_table(
    ['STT', 'Màn hình', 'Loại màn hình', 'Chức năng'],
    [
        ['1', 'Màn hình Đăng ký Hồ sơ VĐV (RegistrationForm — BM13)',
         'Màn hình nhập liệu',
         'Cho phép VĐV nhập và nộp hồ sơ đăng ký tham dự giải đấu (CCCD, ảnh thẻ, hạng mục, đối tác đánh đôi).'],
        ['2', 'Màn hình Duyệt Hồ sơ (RegistrationHubView & ApprovalModal — BM14)',
         'Màn hình nhập liệu / tra cứu',
         'Tiếp nhận, tra cứu và xét duyệt (phê duyệt / từ chối / yêu cầu bổ sung) hồ sơ VĐV.'],
        ['3', 'Màn hình Phân công Trọng tài (MatchAssignmentModal — BM16)',
         'Màn hình nhập liệu',
         'Phân công trọng tài cho từng trận đấu kèm vai trò, trạng thái và cảnh báo ràng buộc.'],
        ['4', 'Màn hình Điều phối lịch đấu (ScheduleView — BM18)',
         'Màn hình điều phối (lịch biểu)',
         'Xếp và điều phối lịch thi đấu theo sân và khung giờ bằng kéo-thả (drag & drop).'],
    ],
    widths=[0.5, 2.6, 1.5, 2.4]
)
doc.add_page_break()

# ============ BM13 ============
h_title('1. BM13 — Màn hình Đăng ký Hồ sơ VĐV (RegistrationForm)')
h_sub('Mô tả màn hình')
body('Màn hình cho phép vận động viên điền và nộp hồ sơ đăng ký tham dự giải đấu. '
     'Họ tên và email được tự động điền từ tài khoản đăng nhập (chỉ đọc). VĐV nhập số CCCD '
     '(12 chữ số), tải lên ảnh thẻ 3x4 (JPG/PNG, tối đa 5MB) và chọn hạng mục thi đấu. '
     'Nếu hạng mục là nội dung đánh đôi (doubles / mixed_doubles), hệ thống hiện khối chọn '
     'VĐV đối tác bắt buộc. Dữ liệu được kiểm tra bằng Zod schema trước khi nộp.')
img_placeholder('Hình 1. Màn hình Đăng ký Hồ sơ VĐV (BM13)')

h_sub('Mô tả các đối tượng trên màn hình')
add_table(
    ['STT', 'Tên', 'Kiểu', 'Ràng buộc', 'Chức năng'],
    [
        ['1', 'Họ và tên VĐV', 'TextBox (disabled)', 'Chỉ đọc, tự động điền', 'Hiển thị tên VĐV từ tài khoản đăng nhập.'],
        ['2', 'Email nhận xác nhận', 'TextBox (disabled)', 'Chỉ đọc, tự động điền', 'Hiển thị email nhận thông báo xác nhận.'],
        ['3', 'Số Căn cước công dân', 'TextBox', 'Bắt buộc; đúng 12 chữ số (regex \\d{12})', 'Nhập số CCCD của VĐV.'],
        ['4', 'Ảnh thẻ 3x4', 'File upload', 'Bắt buộc; chỉ JPG/PNG; ≤ 5MB', 'Tải lên và xem trước ảnh thẻ.'],
        ['5', 'Hạng mục thi đấu', 'ComboBox (select)', 'Bắt buộc chọn 1 hạng mục', 'Chọn nội dung thi đấu đăng ký.'],
        ['6', 'Chọn VĐV đối tác', 'ComboBox (select)', 'Bắt buộc khi là nội dung đánh đôi', 'Chọn VĐV đánh cặp cho nội dung đôi.'],
        ['7', 'Nút Huỷ bỏ', 'Button', '—', 'Đóng form, huỷ thao tác đăng ký.'],
        ['8', 'Nút Nộp hồ sơ', 'Button (submit)', 'Disabled khi đang gửi', 'Gửi hồ sơ đăng ký xuống hệ thống.'],
    ],
    widths=[0.4, 1.6, 1.3, 1.8, 2.4]
)
h_sub('Danh sách biến cố và xử lý tương ứng')
add_table(
    ['STT', 'Biến cố', 'Xử lý'],
    [
        ['1', 'Nhập CCCD sai định dạng', 'Hiện lỗi "CCCD phải gồm đúng 12 chữ số" dưới ô nhập, viền ô chuyển màu cảnh báo.'],
        ['2', 'Tải ảnh sai định dạng (không JPG/PNG)', 'Hiện alert "Chỉ hỗ trợ ảnh JPG hoặc PNG" và xoá file đã chọn.'],
        ['3', 'Tải ảnh vượt quá 5MB', 'Hiện alert "Kích thước ảnh không vượt quá 5MB" và xoá file đã chọn.'],
        ['4', 'Chọn hạng mục đánh đôi', 'Hiện khối cảnh báo yêu cầu chọn VĐV đối tác (bắt buộc).'],
        ['5', 'Nộp khi đánh đôi chưa chọn đối tác', 'Hiện alert "Vui lòng chọn VĐV đối tác cho nội dung đánh đôi!" và chặn gửi.'],
        ['6', 'Chọn nút Nộp hồ sơ (hợp lệ)', 'Gọi onSubmit, gửi dữ liệu hồ sơ xuống CSDL, nút hiển thị "Đang gửi...".'],
        ['7', 'Chọn nút Huỷ bỏ', 'Gọi onCancel, đóng form mà không lưu.'],
    ],
    widths=[0.4, 2.6, 4.0]
)
doc.add_page_break()

# ============ BM14 ============
h_title('2. BM14 — Màn hình Duyệt Hồ sơ (RegistrationHubView & ApprovalModal)')
h_sub('Mô tả màn hình')
body('Màn hình trung tâm cho Ban tổ chức tiếp nhận và xét duyệt hồ sơ. Gồm 4 tab: Tiếp nhận '
     'đăng ký trực tuyến (BM9), Tiếp nhận hồ sơ offline (BM8), Tra cứu/Điều chỉnh hồ sơ VĐV (BM15) '
     'và Theo nội dung thi đấu. Có thanh tìm kiếm (theo tên VĐV / CLB / mã hồ sơ) và bộ lọc theo '
     'nội dung. Bảng danh sách hiện mã HS, VĐV, đoàn/CLB, hạng mục và trạng thái. Nhấn nút "Duyệt" '
     'mở hộp thoại ApprovalModal để phê duyệt, từ chối hoặc yêu cầu bổ sung kèm lý do.')
img_placeholder('Hình 2a. Màn hình tiếp nhận & danh sách hồ sơ (RegistrationHubView)')
img_placeholder('Hình 2b. Hộp thoại duyệt hồ sơ (ApprovalModal — BM14)')

h_sub('Mô tả các đối tượng trên màn hình')
add_table(
    ['STT', 'Tên', 'Kiểu', 'Ràng buộc', 'Chức năng'],
    [
        ['1', 'Tab tiếp nhận/tra cứu', 'Tab (4 mục)', 'Chọn 1 tab', 'Chuyển giữa online/offline/tra cứu/theo nội dung.'],
        ['2', 'Ô tìm kiếm', 'TextBox', '—', 'Lọc hồ sơ theo tên VĐV, CLB hoặc mã HS.'],
        ['3', 'Bộ lọc nội dung', 'ComboBox', '—', 'Lọc theo hạng mục (MS/WS/MD/WD/XD).'],
        ['4', 'Bảng danh sách hồ sơ', 'Table (grid)', 'Chỉ đọc', 'Hiển thị mã HS, VĐV, CLB, hạng mục, trạng thái.'],
        ['5', 'Nhãn trạng thái', 'Badge', '—', 'Hiển thị trạng thái (Chờ duyệt / approved...).'],
        ['6', 'Nút Duyệt', 'Button', '—', 'Mở hộp thoại ApprovalModal cho hồ sơ.'],
        ['7', 'Khối quyết định duyệt', 'Button group', 'Bắt buộc chọn 1', 'Chọn: Phê duyệt / Yêu cầu bổ sung / Từ chối.'],
        ['8', 'Ô Lý do', 'TextArea', 'Bắt buộc khi từ chối / yêu cầu bổ sung', 'Nhập lý do, gửi kèm email cho VĐV.'],
        ['9', 'Nút Xác nhận duyệt', 'Button (submit)', 'Disabled khi thiếu quyết định/lý do', 'Lưu quyết định duyệt xuống CSDL.'],
        ['10', 'Nút Huỷ / Đóng', 'Button', '—', 'Đóng hộp thoại mà không lưu.'],
    ],
    widths=[0.4, 1.6, 1.3, 1.8, 2.4]
)
h_sub('Danh sách biến cố và xử lý tương ứng')
add_table(
    ['STT', 'Biến cố', 'Xử lý'],
    [
        ['1', 'Nhập từ khoá tìm kiếm', 'Lọc danh sách hồ sơ theo tên VĐV / CLB / mã HS theo thời gian thực.'],
        ['2', 'Chọn bộ lọc nội dung', 'Lọc danh sách hồ sơ theo category_code đã chọn.'],
        ['3', 'Chọn nút Duyệt trên 1 dòng', 'Mở hộp thoại ApprovalModal với dữ liệu hồ sơ tương ứng.'],
        ['4', 'Chọn "Phê duyệt"', 'Cập nhật trạng thái = approved; ô lý do không bắt buộc.'],
        ['5', 'Chọn "Từ chối" / "Yêu cầu bổ sung" mà chưa nhập lý do', 'Hiện alert "Vui lòng nhập lý do" và chặn gửi; viền ô lý do cảnh báo.'],
        ['6', 'Chọn nút Xác nhận duyệt (hợp lệ)', 'Gọi API updateStatus, hiện toast "Đã phê duyệt hồ sơ thành công!", đóng modal, tải lại danh sách.'],
        ['7', 'API trả về lỗi', 'Hiện toast lỗi với thông điệp từ server.'],
        ['8', 'Danh sách rỗng / đang tải', 'Hiển thị dòng "Đang tải..." hoặc "Không có hồ sơ nào".'],
    ],
    widths=[0.4, 2.8, 3.8]
)
doc.add_page_break()

# ============ BM16 ============
h_title('3. BM16 — Màn hình Phân công Trọng tài (MatchAssignmentModal)')
h_sub('Mô tả màn hình')
body('Hộp thoại phân công trọng tài cho một trận đấu. Hiển thị thông tin trận (tên, thời gian, sân). '
     'Người dùng chọn trọng tài từ danh sách; sau khi chọn, hệ thống hiện chứng chỉ, số trận đã phân '
     'trong ngày (giới hạn 4 trận/ngày) và đơn vị/CLB. Hệ thống cảnh báo khi trọng tài quá tải (≥4 trận), '
     'cùng đơn vị với VĐV, hoặc khi thao tác trong vòng 15 phút trước giờ thi đấu (chặn lưu). '
     'Cho phép chọn vai trò (chính/biên/giao cầu), trạng thái và ghi chú.')
img_placeholder('Hình 3. Hộp thoại Phân công Trọng tài (BM16)')

h_sub('Mô tả các đối tượng trên màn hình')
add_table(
    ['STT', 'Tên', 'Kiểu', 'Ràng buộc', 'Chức năng'],
    [
        ['1', 'Thông tin trận đấu', 'Label (chỉ đọc)', '—', 'Hiển thị tên trận, thời gian, sân.'],
        ['2', 'Chọn Trọng tài', 'ComboBox (select)', 'Bắt buộc (required)', 'Chọn trọng tài theo mã/tên.'],
        ['3', 'Thẻ thông tin trọng tài', 'Info card', 'Hiện khi đã chọn', 'Hiển thị chứng chỉ, số trận/ngày, đơn vị.'],
        ['4', 'Vai trò', 'ComboBox', 'Mặc định trọng tài chính', 'Chọn vai trò: chính / biên / giao cầu.'],
        ['5', 'Trạng thái', 'ComboBox', 'Mặc định pending', 'Chọn: Chờ xác nhận / Đã duyệt / Huỷ.'],
        ['6', 'Ghi chú', 'TextArea', '—', 'Nhập ghi chú phân công.'],
        ['7', 'Vùng cảnh báo', 'Alert', '—', 'Hiện cảnh báo quá tải / cùng đơn vị / sát giờ.'],
        ['8', 'Nút Lưu phân công', 'Button (submit)', 'Disabled khi chưa chọn TT hoặc sát giờ (<15p)', 'Lưu phân công trọng tài.'],
        ['9', 'Nút Huỷ', 'Button', '—', 'Đóng hộp thoại mà không lưu.'],
    ],
    widths=[0.4, 1.6, 1.3, 1.9, 2.3]
)
h_sub('Danh sách biến cố và xử lý tương ứng')
add_table(
    ['STT', 'Biến cố', 'Xử lý'],
    [
        ['1', 'Chọn trọng tài', 'Hiện thẻ thông tin (chứng chỉ, số trận/ngày, đơn vị) của trọng tài đó.'],
        ['2', 'Trọng tài đã phân ≥ 4 trận/ngày', 'Hiện cảnh báo "Trọng tài đã đạt giới hạn 4 trận/ngày".'],
        ['3', 'Trọng tài cùng đơn vị với VĐV', 'Hiện cảnh báo xung đột đơn vị (màu đỏ).'],
        ['4', 'Thao tác trong vòng 15 phút trước giờ đấu', 'Hiện cảnh báo "Không thể hoán đổi... trong 15 phút" và vô hiệu nút Lưu.'],
        ['5', 'Chọn nút Lưu phân công (hợp lệ)', 'Gọi onSave với {matchId, refereeId, role, status, notes}.'],
        ['6', 'Chưa chọn trọng tài', 'Nút Lưu bị vô hiệu (disabled, mờ).'],
        ['7', 'Chọn nút Huỷ / đóng (X)', 'Đóng hộp thoại mà không lưu.'],
    ],
    widths=[0.4, 2.8, 3.8]
)
doc.add_page_break()

# ============ BM18 ============
h_title('4. BM18 — Màn hình Điều phối lịch đấu (ScheduleView)')
h_sub('Mô tả màn hình')
body('Màn hình điều phối lịch thi đấu dạng lưới thời gian (timeline). Cột trái là danh sách các '
     'trận chưa xếp lịch; vùng phải là lưới gồm hàng theo sân và cột theo khung giờ (08:00–20:00). '
     'Người dùng kéo-thả (drag & drop) thẻ trận từ sidebar vào ô sân-giờ để xếp lịch, hoặc nhấp vào '
     'thẻ/trận để mở hộp thoại xếp lịch chi tiết (chọn sân, thời gian, thời lượng). Có thanh chọn ngày '
     '(theo khoảng ngày giải đấu), bộ lọc trạng thái và hạng mục. Sân đang bảo trì hiện gạch chéo và '
     'không cho xếp lịch.')
img_placeholder('Hình 4a. Màn hình Điều phối lịch đấu — lưới timeline (BM18)')
img_placeholder('Hình 4b. Hộp thoại Xếp lịch trận chi tiết')

h_sub('Mô tả các đối tượng trên màn hình')
add_table(
    ['STT', 'Tên', 'Kiểu', 'Ràng buộc', 'Chức năng'],
    [
        ['1', 'Bộ lọc trạng thái', 'ComboBox', '—', 'Lọc trận theo trạng thái (sắp/đang/hoàn thành).'],
        ['2', 'Bộ lọc hạng mục', 'ComboBox', '—', 'Lọc trận theo hạng mục thi đấu.'],
        ['3', 'Thanh chọn ngày', 'Button group (tabs)', 'Chọn 1 ngày', 'Chọn ngày thi đấu trong khoảng giải.'],
        ['4', 'Sidebar trận chưa xếp lịch', 'List (draggable)', '—', 'Hiển thị các trận chưa có sân; kéo-thả để xếp.'],
        ['5', 'Thẻ trận (card)', 'Draggable card', '—', 'Kéo-thả hoặc nhấp để xếp lịch trận.'],
        ['6', 'Lưới timeline (sân x giờ)', 'Grid / drop zone', '—', 'Nhận thả trận vào ô sân-giờ.'],
        ['7', 'Khối trận đã xếp', 'Block', 'Không hiện trận đã huỷ', 'Hiển thị trận trên lưới theo giờ/sân/thời lượng.'],
        ['8', 'Chọn Sân', 'ComboBox (modal)', 'Bắt buộc', 'Chọn sân trong hộp thoại xếp lịch.'],
        ['9', 'Thời gian', 'datetime-local', 'Bắt buộc', 'Chọn ngày-giờ thi đấu.'],
        ['10', 'Thời lượng (phút)', 'NumberBox', 'Bước 15, tối thiểu 15', 'Nhập thời lượng dự kiến trận.'],
        ['11', 'Nút Lưu lịch', 'Button', 'Cần chọn sân & thời gian', 'Lưu lịch trận xuống CSDL.'],
        ['12', 'Nút Huỷ', 'Button', '—', 'Đóng hộp thoại xếp lịch.'],
    ],
    widths=[0.4, 1.7, 1.4, 1.5, 2.5]
)
h_sub('Danh sách biến cố và xử lý tương ứng')
add_table(
    ['STT', 'Biến cố', 'Xử lý'],
    [
        ['1', 'Chọn ngày trên thanh ngày', 'Lọc và hiển thị các trận đã xếp của ngày đó trên lưới.'],
        ['2', 'Kéo thẻ trận từ sidebar thả vào ô sân-giờ', 'Tính ngày-giờ từ ô thả, gọi API scheduleMatch, toast "Cập nhật lịch thi đấu thành công", tải lại.'],
        ['3', 'Kéo qua ô (dragover/dragleave)', 'Thêm/bỏ lớp "drag-over" làm nổi bật ô đích.'],
        ['4', 'Nhấp vào thẻ/trận', 'Mở hộp thoại xếp lịch với sân/thời gian/thời lượng của trận.'],
        ['5', 'Chọn nút Lưu lịch (hợp lệ)', 'Gọi scheduleMatch với {courtId, scheduledAt, estimatedDurationMins}, toast thành công, đóng modal, tải lại.'],
        ['6', 'API xếp lịch lỗi (vd trùng sân/giờ)', 'Hiện toast lỗi "Lỗi xếp lịch" hoặc thông điệp từ server.'],
        ['7', 'Sân đang bảo trì', 'Hiển thị lớp gạch chéo "Bảo trì — không xếp lịch".'],
        ['8', 'Không có trận chưa xếp', 'Hiển thị "Không có trận trống" trong sidebar.'],
    ],
    widths=[0.4, 2.8, 3.8]
)

doc.save(r'C:\Users\FPT\CNPM-BADMINTON-TOURNAMENT\BaoCao_UI_ThamGia_PhanCongLichDau.docx')
print('OK saved')
